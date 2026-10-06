from pathlib import Path

from docx import Document
from lxml import etree
from pypdf import PdfReader
from pptx import Presentation

def load_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)

def load_docx(path: Path) -> str:
    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)

def load_pptx(path: Path) -> str:
    presentation = Presentation(str(path))
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.text
    )

def load_xml(path: Path) -> str:
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    tree = etree.parse(str(path), parser=parser)
    return " ".join(tree.getroot().itertext())

def load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8")
